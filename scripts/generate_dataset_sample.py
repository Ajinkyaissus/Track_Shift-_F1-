"""
scripts/generate_dataset_sample.py — Generates comprehensive DATASET_SAMPLE.md and docs/TrackShift_Dataset_Sample.docx.
Covers:
1. laps.parquet (Telemetry & Behavioral Features)
2. residual_ledger.parquet (Lap Time Loss & Cumulative Tyre Debt)
3. circuit_geometry.parquet (24-Circuit GPS Track Geometry)
4. bootstrap_uncertainty.parquet (Cluster Bootstrap Confidence Bounds)
5. baseline_predictions.parquet (Stage 1 Linear M1 Model Predictions)
6. tyredebt.db (SQLite Relational Schema)
"""

import os
import sys
import pandas as pd
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
DOCS_DIR = os.path.join(BASE_DIR, "docs")
MD_OUT = os.path.join(BASE_DIR, "DATASET_SAMPLE.md")
DOCX_OUT = os.path.join(DOCS_DIR, "TrackShift_Dataset_Sample.docx")


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_markdown_report():
    laps_df = pd.read_parquet(os.path.join(DATA_DIR, "laps.parquet"))
    ledger_df = pd.read_parquet(os.path.join(DATA_DIR, "residual_ledger.parquet"))
    geom_df = pd.read_parquet(os.path.join(DATA_DIR, "circuit_geometry.parquet"))
    boot_df = pd.read_parquet(os.path.join(DATA_DIR, "bootstrap_uncertainty.parquet"))
    base_df = pd.read_parquet(os.path.join(DATA_DIR, "baseline_predictions.parquet"))

    md = f"""# TrackShift — Formula 1 Real Telemetry & AI/ML Dataset Sample Reference

> **Platform:** TrackShift Formula 1 Tyre Debt Intelligence Platform  
> **Telemetry Source:** Official FastF1 Historical Ingestion (2024 & 2025 FIA Formula 1 World Championship)  
> **Total Records Audited:** 18,513 telemetry laps, 16,376 debt ledger observations, 15,789 circuit geometry coordinates  
> **Format:** Apache Parquet (Columnar snappy-compressed) + SQLite Relational Schema (`tyredebt.db`)

---

## Table of Contents
1. [Dataset Overview & Summary Statistics](#1-dataset-overview--summary-statistics)
2. [Dataset 1: `laps.parquet` — Telemetry & Behavioral Features](#2-dataset-1-lapsparquet--telemetry--behavioral-features)
3. [Dataset 2: `residual_ledger.parquet` — Contextual Residuals & Cumulative Tyre Debt](#3-dataset-2-residual_ledgerparquet--contextual-residuals--cumulative-tyre-debt)
4. [Dataset 3: `circuit_geometry.parquet` — 24-Circuit GPS Track Coordinates & Speeds](#4-dataset-3-circuit_geometryparquet--24-circuit-gps-track-coordinates--speeds)
5. [Dataset 4: `bootstrap_uncertainty.parquet` — 1,000-Iteration Cluster Bootstrap Intervals](#5-dataset-4-bootstrap_uncertaintyparquet--1000-iteration-cluster-bootstrap-intervals)
6. [Dataset 5: `baseline_predictions.parquet` — Stage 1 Linear M1 Lap Time Loss Baseline](#6-dataset-5-baseline_predictionsparquet--stage-1-linear-m1-lap-time-loss-baseline)
7. [Relational SQLite Schema (`api/tyredebt.db`)](#7-relational-sqlite-schema-apityredebtdb)

---

## 1. Dataset Overview & Summary Statistics

| Dataset Name | File Path | Records | File Size | Primary Key / Indexing | Domain Scope |
| :--- | :--- | :---: | :---: | :--- | :--- |
| **Telemetry Laps** | `data/laps.parquet` | 18,513 | ~680 KB | `stint_id`, `lap_number` | Multi-season laps, driving metrics, fuel loads |
| **Residual Debt Ledger** | `data/residual_ledger.parquet` | 16,376 | ~213 KB | `stint_id`, `lap_number` | Baseline lap loss, residual error, cumulative tyre debt |
| **Circuit Geometry** | `data/circuit_geometry.parquet` | 15,789 | ~798 KB | `circuit_id`, `point_order` | 24 Grand Prix tracks, X/Y GPS coordinates, speeds |
| **Bootstrap Uncertainty** | `data/bootstrap_uncertainty.parquet` | 131,670 | ~715 KB | `stint_id`, `feature`, `delta_pct` | 95% Confidence bounds [CI lower, CI upper, margin] |
| **Baseline Predictions** | `data/baseline_predictions.parquet` | 16,376 | ~64 KB | `stint_id`, `lap_number` | Parsimonious Stage 1 linear model predictions |
| **Relational SQLite DB** | `api/tyredebt.db` | 10 Tables | ~1.05 MB | Foreign Keys (`session_id`, `driver_id`, `circuit_id`) | Full normalized relational database for FastF1 |

---

## 2. Dataset 1: `laps.parquet` — Telemetry & Behavioral Features

Contains lap-by-lap high-frequency extracted metrics from FastF1 car telemetry, throttle/brake telemetry, and fuel estimates across all 2024 and 2025 race weekends.

### Schema & Data Types
- `stint_id` (*string*): Composite identifier format `{{year}}_{{circuit}}_{{session}}_{{driver}}_{{stint_num}}` (e.g. `2024_bahrain_R_VER_1`)
- `circuit_id` (*string*): Standardized slug (e.g. `bahrain`, `monza`, `silverstone`, `spa`)
- `session_id` (*string*): Session slug (e.g. `2024_bahrain_R`)
- `driver_id` (*string*): 3-letter driver code (`VER`, `HAM`, `NOR`, `LEC`, `PIA`, `SAI`, etc.)
- `compound` (*string*): Tyre compound name (`SOFT`, `MEDIUM`, `HARD`, `INTERMEDIATE`, `WET`)
- `lap_number` (*int64*): Lap number within race
- `lap_time` (*float64*): Lap time in seconds
- `is_green_flag` (*int64*): Binary flag (1 = Green Flag racing, 0 = SC/VSC/In-lap/Out-lap)
- `fuel_load_est` (*float64*): Estimated fuel on board in kilograms (depleting ~1.75 kg/lap from 110 kg initial)
- `braking_aggression` (*float64*): Mean peak deceleration gradient (bar/s or G/s)
- `throttle_transient_smoothness` (*float64*): Standard deviation of throttle application rate (lower = smoother)
- `lateral_dynamics_proxy` (*float64*): High-speed cornering lateral acceleration load proxy
- `kerb_usage` (*float64*): Integrated high-frequency oscillation magnitude over kerbs
- `lockup_flag_rate` (*float64*): Proportion of braking zones exhibiting front tyre rotational lockup

### Sample Records (Top 5 Rows)
```json
{laps_df.head(5).to_json(orient='records', indent=2)}
```

### Formatted Sample Table
| stint_id | circuit_id | driver_id | compound | lap | lap_time (s) | fuel (kg) | braking_aggression | throttle_smoothness | lateral_proxy | kerb_usage | lockup_rate |
| :--- | :--- | :---: | :---: | :---: | :---: | :---: | :---: | :---: | :---: | :---: | :---: |
"""
    for _, r in laps_df.head(5).iterrows():
        md += f"| `{r['stint_id']}` | `{r['circuit_id']}` | **{r['driver_id']}** | `{r['compound']}` | {r['lap_number']} | {r['lap_time']:.3f} | {r['fuel_load_est']:.1f} | {r['braking_aggression']:.2f} | {r['throttle_transient_smoothness']:.4f} | {r['lateral_dynamics_proxy']:.4f} | {r['kerb_usage']:.2f} | {r['lockup_flag_rate']:.4f} |\n"

    md += f"""
---

## 3. Dataset 2: `residual_ledger.parquet` — Contextual Residuals & Cumulative Tyre Debt

The fundamental ledger powering TrackShift's Stage 2 Tyre Debt calculations. Decomposes lap time loss into baseline degradation and cumulative unrecoverable thermal/mechanical deficit.

### Schema & Data Types
- `stint_id` (*string*): Target stint identifier
- `lap_number` (*int64*): Lap number within race
- `actual_lap_time_loss` (*float64*): Observed lap time loss relative to fastest stint lap (seconds)
- `predicted_lap_time_loss` (*float64*): Stage 1 baseline expected loss $\\hat{{y}} = 0.1974 + 0.0400 \\times \\text{{tyre\\_age}}$
- `residual` (*float64*): Contextual residual $\\epsilon_i = y_i - \\hat{{y}}_i$
- `cumulative_debt` (*float64*): Integrated positive excess loss $\\sum \\max(0, \\epsilon_i)$ (seconds)
- `model_version` (*string*): Provenance model version string (e.g. `v4_m1_production_2026-09-10`)

### Sample Records (Top 5 Rows)
```json
{ledger_df.head(5).to_json(orient='records', indent=2)}
```

### Formatted Sample Table
| stint_id | lap_number | actual_loss (s) | predicted_loss (s) | residual (s) | cumulative_debt (s) | model_version |
| :--- | :---: | :---: | :---: | :---: | :---: | :--- |
"""
    for _, r in ledger_df.head(5).iterrows():
        md += f"| `{r['stint_id']}` | {r['lap_number']} | {r['actual_lap_time_loss']:.3f} | {r['predicted_lap_time_loss']:.3f} | {r['residual']:.3f} | **{r['cumulative_debt']:.3f}** | `{r['model_version']}` |\n"

    md += f"""
---

## 4. Dataset 3: `circuit_geometry.parquet` — 24-Circuit GPS Track Coordinates & Speeds

Normalized 2D/3D track layout telemetry for all 24 official Formula 1 circuits. Used for real-time SVG track rendering, corner identification, driver positioning, and speed heatmaps.

### Schema & Data Types
- `X` (*float64*): Raw GPS easting coordinate in decimeters / meters
- `Y` (*float64*): Raw GPS northing coordinate in decimeters / meters
- `Distance` (*float64*): Distance along track centerline from start/finish line (meters)
- `Speed` (*float64*): Typical reference apex/straight speed (km/h)
- `Throttle` (*float64*): Throttle application percentage (0.0 to 100.0)
- `Brake` (*bool*): Binary brake activation flag (`True` / `False`)
- `circuit_id` (*string*): Standardized circuit slug (e.g. `bahrain`, `monza`, `silverstone`)
- `point_order` (*int64*): Sequential index of coordinate points along track spline
- `x_rot` (*float64*): Rotated and centered SVG X coordinate
- `y_rot` (*float64*): Rotated and centered SVG Y coordinate

### Sample Records (Top 5 Rows)
```json
{geom_df.head(5).to_json(orient='records', indent=2)}
```

### Formatted Sample Table
| circuit_id | point_order | Distance (m) | Speed (km/h) | Throttle (%) | Brake | X (GPS) | Y (GPS) | x_rot (SVG) | y_rot (SVG) |
| :--- | :---: | :---: | :---: | :---: | :---: | :---: | :---: | :---: | :---: |
"""
    for _, r in geom_df.head(5).iterrows():
        md += f"| `{r['circuit_id']}` | {r['point_order']} | {r['Distance']:.1f} | {r['Speed']:.1f} | {r['Throttle']:.0f}% | `{r['Brake']}` | {r['X']:.2f} | {r['Y']:.2f} | {r['x_rot']:.2f} | {r['y_rot']:.2f} |\n"

    md += f"""
---

## 5. Dataset 4: `bootstrap_uncertainty.parquet` — 1,000-Iteration Cluster Bootstrap Intervals

Precomputed statistical uncertainty intervals across counterfactual behavioural adjustments ($\\Delta = -50\\%$ to $+50\\%$) using cluster-level resampling (preserving intra-stint correlation).

### Schema & Data Types
- `stint_id` (*string*): Target stint identifier
- `feature` (*string*): Behavioral feature modified (`braking_aggression`, `throttle_transient_smoothness`, `lateral_dynamics_proxy`, `kerb_usage`, `lockup_flag_rate`)
- `delta_pct` (*float64*): Percentage change in behavioral driver input (e.g. `-20.0%`, `+15.0%`)
- `recovered_p50` (*float64*): Median expected tyre debt recovery (seconds)
- `ci_lower` (*float64*): 2.5th empirical percentile of recovery (seconds)
- `ci_upper` (*float64*): 97.5th empirical percentile of recovery (seconds)
- `ci_margin` (*float64*): Half-width uncertainty interval $(CI_{{upper}} - CI_{{lower}}) / 2$
- `is_saturated` (*bool*): True if hypothetical change breaches physical grip limits
- `max_physical_bound` (*float64*): Maximum physically attainable time recovery bound (seconds)

### Sample Records (Top 5 Rows)
```json
{boot_df.head(5).to_json(orient='records', indent=2)}
```

### Formatted Sample Table
| stint_id | feature | delta_pct | recovered_p50 (s) | 95% CI Lower | 95% CI Upper | CI Margin (±s) | Saturated? | Max Physical (s) |
| :--- | :--- | :---: | :---: | :---: | :---: | :---: | :---: | :---: |
"""
    for _, r in boot_df.head(5).iterrows():
        md += f"| `{r['stint_id']}` | `{r['feature']}` | {r['delta_pct']:+.1f}% | **{r['recovered_p50']:.3f}** | {r['ci_lower']:.3f} | {r['ci_upper']:.3f} | ±{r['ci_margin']:.3f} | `{r['is_saturated']}` | {r['max_physical_bound']:.1f} |\n"

    md += f"""
---

## 6. Dataset 5: `baseline_predictions.parquet` — Stage 1 Linear M1 Lap Time Loss Baseline

Stores the frozen parsimonious baseline predictions for all laps across target sessions.

### Schema & Data Types
- `stint_id` (*string*): Stint identifier
- `lap_number` (*int64*): Lap number
- `actual_lap_time_loss` (*float64*): Actual lap time loss over stint minimum pace
- `predicted_lap_time_loss` (*float64*): Baseline linear model prediction
- `model_version` (*string*): Model version string

### Formatted Sample Table
| stint_id | lap_number | actual_loss (s) | predicted_loss (s) | model_version |
| :--- | :---: | :---: | :---: | :--- |
"""
    for _, r in base_df.head(5).iterrows():
        md += f"| `{r['stint_id']}` | {r['lap_number']} | {r['actual_lap_time_loss']:.3f} | {r['predicted_lap_time_loss']:.3f} | `{r['model_version']}` |\n"

    md += """
---

## 7. Relational SQLite Schema (`api/tyredebt.db`)

The relational layer integrates telemetry with high-level Grand Prix metadata and ML model registries.

### Table Inventory & Purpose
1. `seasons`: Supported calendar seasons (2024, 2025) and status.
2. `events`: Grand Prix weekends with round numbers, dates, locations, countries.
3. `sessions`: Practice, Qualifying, and Race session entries with date/time.
4. `drivers`: Full driver registry (VER, HAM, NOR, LEC, etc.) with team names and car numbers.
5. `circuits`: Circuit metadata (lengths, lap records, corner counts, DRS zones, coordinates).
6. `circuit_geometry`: Normalized coordinate splines for live 2D/3D map tracking.
7. `stints`: Tyre stint metadata (compound, start/end laps, total laps, tyre age).
8. `laps`: Normalized lap times, weather flags, sector splits.
9. `baseline_models`: Registered Stage 1 baseline models, parameters, training split manifests, and RMSE.
10. `stage3_models`: Trained PyTorch TCN architectures, causal convolutional weights, embedding dimensions, and test metrics.
"""

    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write(md)

    print(f"[OK] Generated {MD_OUT}")


def build_docx_report():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("TrackShift — Formula 1 Dataset Sample & Schema Reference")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(225, 6, 0) # F1 Red
    title.paragraph_format.space_after = Pt(4)

    # Subtitle
    sub = doc.add_paragraph()
    sub_run = sub.add_run("Official Telemetry, AI/ML Feature Engineering, and Residual Debt Ledgers")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 100, 110)
    sub.paragraph_format.space_after = Pt(14)

    # Metadata Callout Box
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    set_cell_background(meta_cell, "F4F4F6")
    set_cell_margins(meta_cell, 120, 120, 200, 200)

    meta_p = meta_cell.paragraphs[0]
    meta_p.add_run("Data Source: ").bold = True
    meta_p.add_run("FastF1 Official Historical Telemetry (2024 & 2025 F1 Calendar)\n")
    meta_p.add_run("Telemetry Volume: ").bold = True
    meta_p.add_run("18,513 Laps • 16,376 Debt Records • 15,789 Track Geometry Coordinates\n")
    meta_p.add_run("Primary Storage: ").bold = True
    meta_p.add_run("Apache Parquet (Columnar / Snappy) + SQLite Database (api/tyredebt.db)")
    meta_p.runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Helper to add section headers
    def add_section_header(text):
        h = doc.add_paragraph()
        h_run = h.add_run(text)
        h_run.font.size = Pt(15)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(15, 23, 42)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)

    def add_table_from_df(df, col_names, title_str):
        p = doc.add_paragraph()
        p.add_run(title_str).bold = True
        p.paragraph_format.space_after = Pt(4)

        t = doc.add_table(rows=len(df) + 1, cols=len(col_names))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        hdr_cells = t.rows[0].cells
        for idx, col in enumerate(col_names):
            hdr_cells[idx].text = col
            set_cell_background(hdr_cells[idx], "E10600")
            for p_h in hdr_cells[idx].paragraphs:
                for r_h in p_h.runs:
                    r_h.font.bold = True
                    r_h.font.color.rgb = RGBColor(255, 255, 255)
                    r_h.font.size = Pt(8.5)
            set_cell_margins(hdr_cells[idx], 80, 80, 100, 100)

        # Rows
        for r_idx, row in df.iterrows():
            row_cells = t.rows[r_idx + 1].cells
            bg_hex = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row):
                row_cells[c_idx].text = str(val)
                set_cell_background(row_cells[c_idx], bg_hex)
                for p_r in row_cells[c_idx].paragraphs:
                    for r_r in p_r.runs:
                        r_r.font.size = Pt(8.0)
                set_cell_margins(row_cells[c_idx], 60, 60, 100, 100)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Dataset Overview
    add_section_header("1. Dataset Summary & Inventory")
    summary_data = [
        ["Telemetry Laps", "data/laps.parquet", "18,513", "680 KB", "stint_id, lap_number"],
        ["Residual Debt Ledger", "data/residual_ledger.parquet", "16,376", "213 KB", "stint_id, lap_number"],
        ["Circuit Geometry", "data/circuit_geometry.parquet", "15,789", "798 KB", "circuit_id, point_order"],
        ["Bootstrap Uncertainty", "data/bootstrap_uncertainty.parquet", "131,670", "715 KB", "stint_id, feature, delta_pct"],
        ["Baseline Predictions", "data/baseline_predictions.parquet", "16,376", "64 KB", "stint_id, lap_number"]
    ]
    summary_df = pd.DataFrame(summary_data, columns=["Dataset", "File Path", "Records", "Size", "Keys"])
    add_table_from_df(summary_df, summary_df.columns, "Core Columnar Parquet Files:")

    # 2. laps.parquet
    add_section_header("2. Dataset Sample: laps.parquet (Telemetry & Behavioral Heads)")
    laps_df = pd.read_parquet(os.path.join(DATA_DIR, "laps.parquet"))
    sample_laps = laps_df[['stint_id', 'circuit_id', 'driver_id', 'compound', 'lap_number', 'lap_time', 'fuel_load_est', 'braking_aggression', 'throttle_transient_smoothness']].head(5).copy()
    sample_laps['lap_time'] = sample_laps['lap_time'].round(3)
    sample_laps['braking_aggression'] = sample_laps['braking_aggression'].round(2)
    sample_laps['throttle_transient_smoothness'] = sample_laps['throttle_transient_smoothness'].round(4)
    add_table_from_df(sample_laps.reset_index(drop=True), sample_laps.columns, "Sample Rows from laps.parquet:")

    # 3. residual_ledger.parquet
    add_section_header("3. Dataset Sample: residual_ledger.parquet (Stage 2 Tyre Debt)")
    ledger_df = pd.read_parquet(os.path.join(DATA_DIR, "residual_ledger.parquet"))
    sample_ledger = ledger_df[['stint_id', 'lap_number', 'actual_lap_time_loss', 'predicted_lap_time_loss', 'residual', 'cumulative_debt', 'model_version']].head(5).copy()
    sample_ledger['actual_lap_time_loss'] = sample_ledger['actual_lap_time_loss'].round(3)
    sample_ledger['predicted_lap_time_loss'] = sample_ledger['predicted_lap_time_loss'].round(3)
    sample_ledger['residual'] = sample_ledger['residual'].round(3)
    sample_ledger['cumulative_debt'] = sample_ledger['cumulative_debt'].round(3)
    add_table_from_df(sample_ledger.reset_index(drop=True), sample_ledger.columns, "Sample Rows from residual_ledger.parquet:")

    # 4. circuit_geometry.parquet
    add_section_header("4. Dataset Sample: circuit_geometry.parquet (24-Circuit GPS Splines)")
    geom_df = pd.read_parquet(os.path.join(DATA_DIR, "circuit_geometry.parquet"))
    sample_geom = geom_df[['circuit_id', 'point_order', 'Distance', 'Speed', 'Throttle', 'Brake', 'X', 'Y', 'x_rot', 'y_rot']].head(5).copy()
    sample_geom['Distance'] = sample_geom['Distance'].round(1)
    sample_geom['Speed'] = sample_geom['Speed'].round(1)
    sample_geom['X'] = sample_geom['X'].round(1)
    sample_geom['Y'] = sample_geom['Y'].round(1)
    sample_geom['x_rot'] = sample_geom['x_rot'].round(1)
    sample_geom['y_rot'] = sample_geom['y_rot'].round(1)
    add_table_from_df(sample_geom.reset_index(drop=True), sample_geom.columns, "Sample Rows from circuit_geometry.parquet:")

    # 5. bootstrap_uncertainty.parquet
    add_section_header("5. Dataset Sample: bootstrap_uncertainty.parquet (Cluster Bootstrap CIs)")
    boot_df = pd.read_parquet(os.path.join(DATA_DIR, "bootstrap_uncertainty.parquet"))
    sample_boot = boot_df[['stint_id', 'feature', 'delta_pct', 'recovered_p50', 'ci_lower', 'ci_upper', 'ci_margin', 'is_saturated']].head(5).copy()
    sample_boot['recovered_p50'] = sample_boot['recovered_p50'].round(3)
    sample_boot['ci_lower'] = sample_boot['ci_lower'].round(3)
    sample_boot['ci_upper'] = sample_boot['ci_upper'].round(3)
    sample_boot['ci_margin'] = sample_boot['ci_margin'].round(3)
    add_table_from_df(sample_boot.reset_index(drop=True), sample_boot.columns, "Sample Rows from bootstrap_uncertainty.parquet:")

    doc.save(DOCX_OUT)
    print(f"[OK] Generated {DOCX_OUT}")


if __name__ == "__main__":
    build_markdown_report()
    build_docx_report()
